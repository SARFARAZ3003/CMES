"""
PROD.xlsx (HISTORY sheet) -> SEEDHA SQL Server mein insert (optional, fast)
--------------------------------------------------------------------------
Ye script Excel padh ke directly dbo.MPI_COB_T_SERIAL_NO_HISTORY mein
rows daal deta hai - .sql file banane ki zarurat nahi.

Pehle ek baar install karo:
    python -m pip install openpyxl pyodbc

Phir run:
    python load_to_sql.py

Note: ODBC Driver chahiye. SQL Server ke saath aam taur pe
"ODBC Driver 17 for SQL Server" hota hai. Niche DRIVER badal sakte ho.
"""

import openpyxl
import pyodbc

# ---- Config (apne hisaab se badlo) ----
EXCEL_PATH = r"C:\Users\Sarfaraz\Downloads\PROD.xlsx"
SHEET_NAME = "HISTORY"
SERVER = "localhost"          # SSMS mein jo server dikhta hai
DATABASE = "FlexNet"          # jis DB mein table hai (USE wala naam)
DRIVER = "ODBC Driver 17 for SQL Server"
TABLE_NAME = "dbo.MPI_COB_T_SERIAL_NO_HISTORY"
BATCH_SIZE = 1000

COLUMNS = [
    "ID", "PRODUCTID", "SERIALNO", "LOTNO", "WORKORDERNO", "WORKSTATION",
    "STATUS", "PREVIOUSSTATUS", "ENGINEBUILDPROPERTY", "LOCATION", "REPAIRGROUP",
    "REINTRODUCEFLAG", "REINTRODUCEWORKSTATION", "LIFTOFFREASON", "APPLICATION",
    "ISLIFTOFFWHENPAUSED", "LIFTOFFWORKSTATION", "REFERENCEID", "LASTUPDATEON",
    "LASTUPDATEDBY", "CREATEDON", "CREATEDBY", "ACTIVE", "LASTDELETEON",
    "LASTDELETEDBY", "LASTREACTIVATEON", "LASTREACTIVATEDBY", "ARCHIVED",
    "LASTARCHIVEON", "LASTARCHIVEDBY", "LASTRESTOREON", "LASTRESTOREDBY",
    "ROWVERSIONSTAMP",
]

NUM_IDX = {0, 1, 6, 7, 8, 11, 15, 17, 22, 27, 32}
DT_IDX  = {18, 20, 23, 25, 28, 30}


def clean(idx, v):
    # Empty string -> None (NULL). Numeric ke liye int/float, baaki as-is.
    if v is None:
        return None
    if isinstance(v, str) and v.strip() == "":
        return None
    if idx in NUM_IDX:
        f = float(v)
        return int(f) if f.is_integer() else f
    return v  # datetime aur string pyodbc khud handle karta hai


def main():
    print("Excel read kar rahe hain ...")
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    ws = wb[SHEET_NAME]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    print(f"{len(rows)} rows mile.")

    data = []
    for row in rows:
        data.append(tuple(
            clean(c, row[c] if c < len(row) else None)
            for c in range(len(COLUMNS))
        ))

    conn_str = (
        f"DRIVER={{{DRIVER}}};SERVER={SERVER};DATABASE={DATABASE};"
        f"Trusted_Connection=yes;"
    )
    print("SQL Server se connect ho rahe hain ...")
    conn = pyodbc.connect(conn_str)
    cur = conn.cursor()
    cur.fast_executemany = True

    placeholders = ", ".join(["?"] * len(COLUMNS))
    sql = f"INSERT INTO {TABLE_NAME} ({', '.join(COLUMNS)}) VALUES ({placeholders})"

    inserted = 0
    for i in range(0, len(data), BATCH_SIZE):
        batch = data[i:i + BATCH_SIZE]
        cur.executemany(sql, batch)
        conn.commit()
        inserted += len(batch)
        print(f"  {inserted}/{len(data)} inserted ...")

    cur.close()
    conn.close()
    print(f"Done! {inserted} rows insert ho gaye {TABLE_NAME} mein.")


if __name__ == "__main__":
    main()
